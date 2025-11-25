# core/excel_parser.py
import pandas as pd
from openpyxl import load_workbook
import numpy as np

class ExcelParser:
    def __init__(self):
        self.current_data = None
        self.current_headers = None
        self.current_file = None
    
    def read_excel(self, file_path, sheet_name=0):
        self.current_file = file_path
        try:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            df = self._clean_dataframe(df)
            
            self.current_data = df
            self.current_headers = df.columns.tolist()
            
            # Return tuple compatible with UI tables
            data_rows = [tuple(row) for row in df.values]
            return self.current_headers, data_rows
        except Exception as e:
            print(f"Error reading Excel file: {str(e)}")
            raise
    
    def _clean_dataframe(self, df):
        df = df.dropna(how='all').dropna(axis=1, how='all')
        df.columns = [f"Column {i}" if pd.isna(col) else str(col) for i, col in enumerate(df.columns)]
        
        counts = {}
        new_columns = []
        for col in df.columns:
            if col in counts:
                counts[col] += 1
                new_columns.append(f"{col}_{counts[col]}")
            else:
                counts[col] = 0
                new_columns.append(col)
        df.columns = new_columns
        df = df.fillna('')
        return df
    
    def get_sheet_names(self, file_path):
        try:
            workbook = load_workbook(file_path, read_only=True)
            return workbook.sheetnames
        except Exception as e:
            print(f"Error getting sheet names: {str(e)}")
            raise
    
    def get_json_data(self):
        if self.current_data is None:
            return None
        
        records = self.current_data.to_dict('records')
        for record in records:          
            for key, value in record.items():
                if isinstance(value, np.integer):
                    record[key] = int(value)
                elif isinstance(value, np.floating):
                    record[key] = float(value)
                elif isinstance(value, np.ndarray):
                    record[key] = value.tolist()
        return records
    
    def extract_names(self, json_list):
        names = []
        for json_obj in json_list:
            name = json_obj.get('name')
            if name:
                names.append(name)
        return names

    def save_to_excel(self, file_path, data=None, headers=None):
        try:
            if data is not None:
                if isinstance(data, pd.DataFrame):
                    df = data
                else:
                    df = pd.DataFrame(data, columns=headers or self.current_headers)
            else:
                df = self.current_data
            
            if df is not None:
                df.to_excel(file_path, index=False)
                return True
            else:
                raise ValueError("No data to save")
        except Exception as e:
            print(f"Error saving to Excel: {str(e)}")
            raise
